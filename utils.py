import os
import io
import logging
import datetime
from pathlib import Path

from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from pypdf import PdfReader, PdfWriter
import openpyxl
from openpyxl.styles import Font


def setup_logging(log_path):
    Path(log_path).parent.mkdir(parents=True, exist_ok=True)

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
        handlers=[
            logging.FileHandler(log_path, encoding="utf-8"),
            logging.StreamHandler()
        ]
    )


def get_pdfs(folder):
    p = Path(folder)
    if not p.exists():
        raise FileNotFoundError(f"Input folder not found: {folder}")
    return [f for f in p.iterdir() if f.is_file() and f.suffix.lower() == ".pdf"]


def check_stamp(stamp_path):
    if not Path(stamp_path).exists():
        raise FileNotFoundError(f"Stamp not found: {stamp_path}")

    ext = Path(stamp_path).suffix.lower()
    if ext not in [".png", ".jpg", ".jpeg", ".pdf"]:
        raise ValueError(f"Stamp format not supported: {ext}")

    if ext in [".png", ".jpg", ".jpeg"]:
        try:
            Image.open(stamp_path).verify()
        except Exception as e:
            raise ValueError(f"Stamp image is corrupted: {e}")


def make_stamp_overlay(stamp_path, width, height):
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(width, height))

    stamp_w = 150
    stamp_h = 75
    x = width - stamp_w - 20
    y = 20

    img = ImageReader(stamp_path)
    c.drawImage(img, x, y, width=stamp_w, height=stamp_h, mask="auto")
    c.save()

    buf.seek(0)
    return PdfReader(buf)


def stamp_pdf(pdf_path, stamp_path, out_folder):
    try:
        reader = PdfReader(str(pdf_path))
    except Exception as e:
        raise ValueError(f"Corrupted PDF: {e}")

    if reader.is_encrypted:
        raise ValueError("PDF is encrypted")

    if len(reader.pages) == 0:
        raise ValueError("PDF has no pages")

    writer = PdfWriter()
    last = len(reader.pages) - 1

    for i, page in enumerate(reader.pages):
        if i == last:
            w = float(page.mediabox.width)
            h = float(page.mediabox.height)
            overlay = make_stamp_overlay(stamp_path, w, h)
            page.merge_page(overlay.pages[0])
        writer.add_page(page)

    out_name = pdf_path.stem + "_stamped.pdf"
    out_path = Path(out_folder) / out_name

    with open(out_path, "wb") as f:
        writer.write(f)

    return str(out_path)


def init_report(report_path):
    Path(report_path).parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stamping Report"

    ws.append(["File Name", "Stamped", "Timestamp", "Remarks"])

    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 35

    wb.save(report_path)


def log_to_report(report_path, filename, stamped, remarks):
    wb = openpyxl.load_workbook(report_path)
    ws = wb.active
    ts = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws.append([filename, stamped, ts, remarks])
    wb.save(report_path)